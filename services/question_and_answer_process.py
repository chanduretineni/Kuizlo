import io
import logging
import base64
import json
import os
import re
from datetime import datetime
from zipfile import ZipFile
from typing import Tuple, List, Optional
from PIL import Image

import fitz  
import docx  
from openai import OpenAI
from fastapi import HTTPException, Form, File, UploadFile, APIRouter
from models.request_models import AnswerResponse
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch
from reportlab.lib.utils import ImageReader

import matplotlib.pyplot as plt
import seaborn as sns
import pandas as pd
from reportlab.platypus import Image as RLImage  # Renamed to avoid conflict with PIL.Image

from config import OPENAI_API_KEY

# Initialize FastAPI router
router = APIRouter()

# Initialize OpenAI client
client = OpenAI(api_key=OPENAI_API_KEY)

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def encode_image_bytes(image_bytes: bytes) -> str:
    """
    Encodes image bytes into a Base64 string.
    """
    return base64.b64encode(image_bytes).decode("utf-8")


def answer_question(content: str, image: bytes = None, images: list = None) -> str:
    """
    Sends a request to the OpenAI API with the given text content and optional image(s).
    Expects the assistant to return a JSON response with Python plotting code.
    """
    user_payload = {"content": content, "images": []}
    
    if image:
        user_payload["images"].append(f"data:image/jpeg;base64,{encode_image_bytes(image)}")
    if images:
        for img in images:
            user_payload["images"].append(f"data:image/jpeg;base64,{encode_image_bytes(img)}")
    
    user_content = json.dumps(user_payload)
    
    system_prompt = """You are a helpful assistant. When answering questions that require data visualization:
1. You MUST return a valid JSON object with exactly two keys: 'answer' and 'plots'
2. The 'answer' key should contain text with {{PLOT}} placeholders indicating where plots should appear
3. The 'plots' key should contain an array of Python code strings (without ```python markers)
4. Each plot code must be a complete, self-contained matplotlib script
5. DO NOT use code block markers (```) in your response

Example response format:
{
    "answer": "Here's the data visualization {{PLOT}}. As we can see from the graph...",
    "plots": [
        "import matplotlib.pyplot as plt\\ndata = [1, 2, 3, 4, 5]\\nplt.figure(figsize=(8, 6))\\nplt.plot(data)\\nplt.title('Simple Plot')"
    ]
}"""
    
    payload = {
        "model": "gpt-4o",
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_content},
        ]
    }

    try:
        response = client.chat.completions.create(**payload)
        answer = response.choices[0].message.content
        return answer
    except Exception as e:
        raise Exception(f"OpenAI API error: {str(e)}")


def clean_python_code(code: str) -> str:
    """
    Cleans and normalizes Python plot code by removing escape characters and
    ensuring proper formatting.
    """
    # Remove escape characters
    code = code.replace('\\n', '\n')
    code = code.replace('\\t', '\t')
    
    # Remove code block markers if present
    code = re.sub(r'^```python\s*', '', code)
    code = re.sub(r'```$', '', code)
    
    return code.strip()

def execute_python_plot_code(code: str) -> str:
    """
    Executes matplotlib plot code and saves the resulting plot as an image.
    Handles plot cleanup and proper figure management.
    
    Args:
        code (str): Python code that generates a matplotlib plot
        
    Returns:
        str: Path to the saved plot image
    """
    temp_dir = "temp_plots"
    os.makedirs(temp_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    plot_path = os.path.join(temp_dir, f"python_plot_{timestamp}.png")
    
    # Clean any existing plots
    plt.close('all')
    
    # Clean the code before execution
    cleaned_code = clean_python_code(code)
    
    # Remove plt.show() calls as they interfere with saving
    cleaned_code = re.sub(r'plt\.show\(\)', '', cleaned_code)
    
    # Prepare execution environment with necessary modules
    safe_globals = {
        "__builtins__": __builtins__,
        "plt": plt,
        "base64": base64,
        "BytesIO": io.BytesIO,
        "numpy": __import__('numpy'),
        "pd": __import__('pandas')
    }
    local_env = {}
    
    try:
        # Execute the plot code
        exec(cleaned_code, safe_globals, local_env)
        
        # Handle the plot based on execution results
        if "image_base64" in local_env:
            # If code generated a base64 string, decode and save it
            image_data = base64.b64decode(local_env["image_base64"])
            with open(plot_path, "wb") as f:
                f.write(image_data)
        elif plt.get_fignums():
            # If there's an active plot, save it
            plt.savefig(plot_path, bbox_inches='tight', dpi=300)
            plt.close('all')  # Clean up after saving
        else:
            plt.close('all')  # Ensure cleanup
            raise ValueError("No plot was generated by the provided code.")
            
        return plot_path
        
    except Exception as e:
        plt.close('all')  # Ensure cleanup on error
        logger.error(f"Error executing plot code: {str(e)}\nCode:\n{cleaned_code}")
        raise HTTPException(status_code=500, detail=f"Failed to generate plot: {str(e)}")
    finally:
        # Always ensure plots are closed
        plt.close('all')

def extract_json_from_response(response: str) -> dict:
    """
    Extracts JSON from OpenAI response that might contain markdown or additional text.
    
    Args:
        response (str): Raw response from OpenAI API
        
    Returns:
        dict: Parsed JSON object with 'answer' and 'plots' keys
    """
    # Try to find JSON block in markdown code blocks
    json_match = re.search(r'```json\s*(.*?)\s*```', response, re.DOTALL)
    if json_match:
        try:
            return json.loads(json_match.group(1))
        except json.JSONDecodeError:
            pass

    # Try to find any JSON-like structure in the response
    try:
        # Find text that looks like JSON object
        json_pattern = r'\{[\s\S]*"answer"[\s\S]*"plots"[\s\S]*\}'
        json_match = re.search(json_pattern, response)
        if json_match:
            return json.loads(json_match.group(0))
    except json.JSONDecodeError:
        pass
    
    # If no valid JSON found, try to extract code blocks and create structured response
    code_blocks = re.findall(r'```(?:python)?\s*(.*?)\s*```', response, re.DOTALL)
    if code_blocks:
        # Create a structured response with the text before the first code block as answer
        text_parts = response.split('```')
        answer_text = text_parts[0].strip()
        return {
            "answer": answer_text + " {{PLOT}}",
            "plots": code_blocks
        }
    
    # If no code blocks found, return the raw text as answer
    return {
        "answer": response,
        "plots": []
    }

def create_pdf_from_response(answer: str, question_type: str = None) -> str:
    """
    Creates a PDF document from the API response with proper formatting,
    generating and embedding plots where indicated.
    """
    results_dir = "results"
    os.makedirs(results_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    question_type = question_type or "response"
    filename = f"{question_type}_{timestamp}_answer.pdf"
    filepath = os.path.join(results_dir, filename)
    
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=72,
        leftMargin=72,
        topMargin=72,
        bottomMargin=72
    )
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        spaceAfter=30,
        alignment=1
    )
    content_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontSize=12,
        leading=14,
        spaceAfter=12
    )
    
    story = []
    title = question_type.replace("_", " ").title()
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 12))
    
    try:
        # Parse the response and extract JSON content
        try:
            response_dict = extract_json_from_response(answer)
            answer_text = response_dict.get("answer", "")
            plots_list = response_dict.get("plots", [])
            
            # Clean up the answer text - remove markdown table if present
            answer_text = re.sub(r'\|.*\|[\r\n]*(?:\|.*\|[\r\n]*)*', '', answer_text)
            answer_text = re.sub(r'\s+', ' ', answer_text).strip()
            
        except Exception as e:
            logger.error(f"Failed to parse response: {str(e)}")
            answer_text = answer
            plots_list = []
        
        # Split text by plot placeholders
        segments = answer_text.split("{{PLOT}}")
        
        # Process each segment and add plots
        for i, segment in enumerate(segments):
            if segment.strip():
                story.append(Paragraph(segment.strip(), content_style))
            
            if i < len(plots_list):
                try:
                    plot_path = execute_python_plot_code(plots_list[i])
                    story.append(Spacer(1, 12))
                    story.append(RLImage(plot_path, width=6*inch, height=4*inch))
                    story.append(Spacer(1, 12))
                except Exception as e:
                    logger.error(f"Failed to generate plot {i+1}: {str(e)}")
                    story.append(Paragraph(f"[Failed to generate plot {i+1}]", content_style))
        
        # Build PDF
        doc.build(story)
        
        # Clean up temporary plot files
        temp_dir = "temp_plots"
        if os.path.exists(temp_dir):
            for file in os.listdir(temp_dir):
                os.remove(os.path.join(temp_dir, file))
            os.rmdir(temp_dir)
        
        return filepath
        
    except Exception as e:
        logger.error(f"PDF generation error: {str(e)}")
        story.append(Paragraph(str(answer), content_style))
        doc.build(story)
        return filepath
    
async def extract_text_and_images(file_extension: str, file_io: io.BytesIO) -> Tuple[str, List[bytes]]:
    """
    Extract text and images from PDF or DOCX files.
    Returns a tuple: (extracted_text, list_of_image_bytes)
    """
    extracted_text = ""
    images = []

    if file_extension.lower() == 'pdf':
        try:
            file_io.seek(0)
            doc = fitz.open(stream=file_io, filetype="pdf")
            for page in doc:
                extracted_text += page.get_text() + "\n"
                image_list = page.get_images(full=True)
                for img in image_list:
                    xref = img[0]
                    pix = fitz.Pixmap(doc, xref)
                    if pix.n < 5:  # GRAY or RGB
                        img_bytes = pix.tobytes("png")
                    else:  # CMYK: convert to RGB first.
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                        img_bytes = pix.tobytes("png")
                    images.append(img_bytes)
            return extracted_text, images
        except Exception as e:
            logger.error(f"Error extracting PDF content: {str(e)}")
            raise HTTPException(status_code=500, detail="Error processing PDF file")

    elif file_extension.lower() == 'docx':
        try:
            file_io.seek(0)
            document = docx.Document(file_io)
            extracted_text = "\n".join([para.text for para in document.paragraphs])
            file_io.seek(0)
            with ZipFile(file_io) as docx_zip:
                for file_name in docx_zip.namelist():
                    if file_name.startswith("word/media/"):
                        images.append(docx_zip.read(file_name))
            return extracted_text, images
        except Exception as e:
            logger.error(f"Error extracting DOCX content: {str(e)}")
            raise HTTPException(status_code=500, detail="Error processing DOCX file")
    else:
        raise HTTPException(status_code=400, detail="Unsupported file extension for extraction")


def format_response_to_pdf(response: str, file_prefix: str = None) -> str:
    """
    Wrapper function to handle the PDF generation process.
    
    Args:
        response (str): The API response to format.
        file_prefix (str): Optional prefix for the filename (e.g., 'math', 'science').
    
    Returns:
        str: Path to the generated PDF file.
    """
    try:
        content_type = file_prefix or "general"
        if "\\[" in response or "\\(" in response:
            content_type = "math_problem"
        
        pdf_path = create_pdf_from_response(response, content_type)
        return pdf_path
        
    except Exception as e:
        raise Exception(f"Error generating PDF: {str(e)}")
    

def clean_latex_formatting(text: str) -> str:
    """
    Removes LaTeX formatting from text while preserving line breaks and
    mathematical content.
    
    Args:
        text (str): Text containing LaTeX formatting.
        
    Returns:
        str: Cleaned text with LaTeX formatting removed but structure preserved.
    """
    text = text.replace('\r\n', '\n')
    cleaned_text = re.sub(r'\\\[(.*?)\\\]', r'\1', text, flags=re.DOTALL)
    cleaned_text = re.sub(r'\\\((.*?)\\\)', r'\1', cleaned_text)
    
    latex_commands = [
        (r'\\frac\{(.*?)\}\{(.*?)\}', r'\1/\2'),
        (r'\\sqrt\{(.*?)\}', r'√(\1)'),
        (r'\\cdot', r'×'),
        (r'\\times', r'×'),
        (r'\\div', r'÷'),
        (r'\\pm', r'±'),
        (r'\\leq', r'≤'),
        (r'\\geq', r'≥'),
        (r'\\neq', r'≠'),
        (r'\\alpha', r'α'),
        (r'\\beta', r'β'),
        (r'\\theta', r'θ'),
        (r'\\pi', r'π'),
        (r'\\boxed\{(.*?)\}', r'\1'),
    ]
    
    for pattern, replacement in latex_commands:
        cleaned_text = re.sub(pattern, replacement, cleaned_text)
    
    cleaned_text = cleaned_text.replace('\\', '')
    
    lines = cleaned_text.split('\n')
    cleaned_lines = []
    
    previous_line_empty = False
    for line in lines:
        cleaned_line = re.sub(r' +', ' ', line.strip())
        if cleaned_line:
            cleaned_lines.append(cleaned_line)
            previous_line_empty = False
        elif not previous_line_empty:
            cleaned_lines.append('')
            previous_line_empty = True
    
    return '\n'.join(cleaned_lines)
